"""Excel hisobot servisi."""
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import (
    Alignment,
    Border,
    Font,
    PatternFill,
    Side,
)
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy.orm import selectinload

from src.database.models.truck import Truck
from src.database.models.truck_step import TruckStep
from src.utils.constants import (
    PRIORITY_NAMES,
    STEP_NAMES,
)


# Uslublar
HEADER_FONT = Font(bold=True, color="FFFFFF", size=11)
HEADER_FILL = PatternFill(
    start_color="2F5496", end_color="2F5496", fill_type="solid"
)
HEADER_ALIGN = Alignment(horizontal="center", vertical="center")
BORDER = Border(
    left=Side(style="thin"),
    right=Side(style="thin"),
    top=Side(style="thin"),
    bottom=Side(style="thin"),
)


def _style_header(ws, row: int, num_cols: int) -> None:
    """Header uslubini qo'llash."""
    for col in range(1, num_cols + 1):
        cell = ws.cell(row=row, column=col)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = HEADER_ALIGN
        cell.border = BORDER


def _auto_width(ws) -> None:
    """Ustun kengliklarini avtomatik sozlash."""
    for column in ws.columns:
        max_length = 0
        column_letter = column[0].column_letter
        for cell in column:
            try:
                if cell.value:
                    max_length = max(max_length, len(str(cell.value)))
            except Exception:
                pass
        ws.column_dimensions[column_letter].width = min(max_length + 4, 40)


async def generate_trucks_report(
    session: AsyncSession,
    status_filter: str | None = None,
) -> BytesIO:
    """Trucklar hisoboti (Excel)."""
    stmt = (
        select(Truck)
        .options(selectinload(Truck.steps))
        .order_by(Truck.created_at.desc())
    )

    if status_filter:
        stmt = stmt.where(Truck.status == status_filter)

    result = await session.execute(stmt)
    trucks = list(result.scalars().all())

    wb = Workbook()
    ws = wb.active
    ws.title = "Trucklar"

    headers = [
        "ID",
        "Serial",
        "Model",
        "Buyurtmachi",
        "Prioritet",
        "Joriy step",
        "Holat",
        "Yaratilgan",
        "Tugatilgan",
        "Davomiylik (kun)",
    ]
    ws.append(headers)
    _style_header(ws, 1, len(headers))

    for truck in trucks:
        priority_name = PRIORITY_NAMES.get(truck.priority, truck.priority)

        status_map = {
            "in_progress": "Jarayonda",
            "completed": "Tayyor",
            "cancelled": "Bekor qilingan",
        }
        status_name = status_map.get(truck.status, truck.status)

        duration: int | str = ""
        if truck.completed_at and truck.created_at:
            delta = truck.completed_at - truck.created_at
            duration = delta.days

        ws.append([
            truck.id,
            truck.serial_number,
            truck.model or "—",
            truck.customer or "—",
            priority_name,
            f"{truck.current_step}/6",
            status_name,
            truck.created_at.strftime("%Y-%m-%d %H:%M")
            if truck.created_at
            else "—",
            truck.completed_at.strftime("%Y-%m-%d %H:%M")
            if truck.completed_at
            else "—",
            duration,
        ])

    for row in ws.iter_rows(
        min_row=2, max_row=ws.max_row, max_col=len(headers)
    ):
        for cell in row:
            cell.border = BORDER

    _auto_width(ws)
    ws.freeze_panes = "A2"

    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return output


async def generate_steps_report(
    session: AsyncSession,
    truck_id: int | None = None,
) -> BytesIO:
    """Steplar hisoboti (Excel)."""
    stmt = (
        select(TruckStep)
        .options(
            selectinload(TruckStep.truck),
            selectinload(TruckStep.worker),
            selectinload(TruckStep.qc),
        )
        .order_by(TruckStep.truck_id, TruckStep.step_number)
    )

    if truck_id:
        stmt = stmt.where(TruckStep.truck_id == truck_id)

    result = await session.execute(stmt)
    steps = list(result.scalars().all())

    wb = Workbook()
    ws = wb.active
    ws.title = "Steplar"

    headers = [
        "Truck",
        "Step №",
        "Step nomi",
        "Ishchi",
        "QC",
        "Holat",
        "Izoh",
        "QC izohi",
        "Yuborilgan",
        "Tekshirilgan",
        "Davomiylik (soat)",
    ]
    ws.append(headers)
    _style_header(ws, 1, len(headers))

    for step in steps:
        step_name = STEP_NAMES.get(
            step.step_number, f"Step {step.step_number}"
        )

        status_map = {
            "pending": "Kutilmoqda",
            "in_review": "Tekshiruvda",
            "approved": "Tasdiqlangan",
            "rejected": "Rad etilgan",
        }
        status_name = status_map.get(step.status, step.status)

        duration: float | str = ""
        if step.submitted_at and step.reviewed_at:
            delta = step.reviewed_at - step.submitted_at
            duration = round(delta.total_seconds() / 3600, 1)

        ws.append([
            step.truck.serial_number if step.truck else "—",
            step.step_number,
            step_name,
            step.worker.full_name if step.worker else "—",
            step.qc.full_name if step.qc else "—",
            status_name,
            (step.worker_comment or "—")[:50],
            (step.qc_comment or "—")[:50],
            step.submitted_at.strftime("%Y-%m-%d %H:%M")
            if step.submitted_at
            else "—",
            step.reviewed_at.strftime("%Y-%m-%d %H:%M")
            if step.reviewed_at
            else "—",
            duration,
        ])

    for row in ws.iter_rows(
        min_row=2, max_row=ws.max_row, max_col=len(headers)
    ):
        for cell in row:
            cell.border = BORDER

    _auto_width(ws)
    ws.freeze_panes = "A2"

    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return output
